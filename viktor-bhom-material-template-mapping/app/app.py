from __future__ import annotations

import json
import math
import re
import uuid
from io import BytesIO
from typing import Any
from xml.sax.saxutils import escape
from zipfile import ZIP_DEFLATED, ZipFile

import viktor as vkt
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .catalog import (
    build_template_materials,
    catalogue_source,
    parse_lookup_results,
    record_by_id,
)
from .contracts import (
    ContractError,
    build_handoff,
    material_density,
    parse_mapping_state,
    resolve_inputs,
    value_of,
)
from .parametrization import Parametrization
from .webview import build_mapping_html
from .worker_client import CACHE_VERSION, execute_worker

_DATA_VIEW_ITEM_LIMIT = 100
_MAX_EXTRA_ITEMS_PER_MATERIAL = 12
_MAX_EXTRA_ITEMS_OVERALL = 40
_DYNAMIC_EPD_CONTAINERS = ("EnvironmentalMetrics", "Fragments")
_DYNAMIC_EXCLUDED_KEYS = {
    "_t",
    "BHoM_Guid",
    "Density",
    "Name",
    "Properties",
    "QuantityType",
}
_UNIT_SUFFIXES = (
    ("_kgco2e_per_declared_unit", " kgCO₂e/declared unit"),
    ("_kg_m3", " kg/m³"),
    ("_kg_m2", " kg/m²"),
    ("_kg", " kg"),
    ("_m3", " m³"),
    ("_m2", " m²"),
    ("_m", " m"),
    ("_percent", " %"),
    ("_pct", " %"),
)


_XLSX_SHEETS = ("Material template", "Takeoff", "Workflow handoff")


def _flatten_excel_payload(value: Any, path: str = "") -> list[tuple[str, Any]]:
    if isinstance(value, dict):
        rows: list[tuple[str, Any]] = []
        for key, child in value.items():
            child_path = f"{path}.{key}" if path else str(key)
            rows.extend(_flatten_excel_payload(child, child_path))
        return rows
    if isinstance(value, list):
        rows = []
        for index, child in enumerate(value):
            rows.extend(_flatten_excel_payload(child, f"{path}[{index}]"))
        return rows
    return [(path or "Value", value)]


def _material_template_rows(
    material_template: list[dict[str, Any]],
) -> list[tuple[Any, ...]]:
    rows: list[tuple[Any, ...]] = [("Material", "Field", "Value")]
    for material in material_template:
        material_name = str(material.get("Name") or "Unnamed material")
        rows.extend(
            (material_name, field, value)
            for field, value in _flatten_excel_payload(material)
        )
    return rows


def _takeoff_rows(takeoff: dict[str, Any]) -> list[tuple[Any, ...]]:
    rows: list[tuple[Any, ...]] = [
        (
            "Material",
            "BHoM type",
            "Volume (m³)",
            "Mass (kg)",
            "Area (m²)",
            "Length (m)",
            "Count",
            "Density (kg/m³)",
        )
    ]
    for item in takeoff.get("MaterialTakeoffItems", []):
        material = item.get("Material", {})
        rows.append(
            (
                material.get("Name"),
                material.get("_t"),
                item.get("Volume"),
                item.get("Mass"),
                item.get("Area"),
                item.get("Length"),
                item.get("NumberItem"),
                material.get("Density"),
            )
        )
    return rows


def _handoff_rows(handoff: dict[str, Any]) -> list[tuple[Any, ...]]:
    target = handoff["target"]
    source = handoff["source"]
    catalogue = source["catalogue"]
    params = handoff["params"]
    takeoff = json.loads(params["takeoff_json"])
    template_materials = json.loads(params["template_materials_json"])
    return [
        ("Field", "Value"),
        ("Schema version", handoff["schema_version"]),
        ("Target app", target["app"]),
        ("Target method", target["method_name"]),
        ("Source app", source["app"]),
        ("Catalogue", catalogue["dataset"]),
        ("BHoM component", catalogue["toolkit"]),
        ("Library path", catalogue["library_path"]),
        ("Available datasets", catalogue["dataset_count"]),
        ("Available EPDs", catalogue["epd_count"]),
        ("Gross floor area (m²)", params["gross_floor_area_m2"]),
        ("Modules", ", ".join(params["modules"])),
        ("Prioritise template materials", params["prioritise_template_materials"]),
        ("Takeoff type", takeoff["_t"]),
        ("Takeoff name", takeoff.get("Name")),
        ("Takeoff items", len(takeoff["MaterialTakeoffItems"])),
        ("Template materials", len(template_materials)),
    ]


def _xlsx_text(value: Any) -> str:
    text = str(value)
    return "".join(
        character for character in text if character in "\t\n\r" or ord(character) >= 32
    )


def _xlsx_cell(reference: str, value: Any, style: int) -> str:
    if isinstance(value, bool):
        return f'<c r="{reference}" s="{style}" t="b"><v>{int(value)}</v></c>'
    if (
        isinstance(value, (int, float))
        and not isinstance(value, bool)
        and (not isinstance(value, float) or math.isfinite(value))
    ):
        return f'<c r="{reference}" s="{style}"><v>{value}</v></c>'
    text = escape(_xlsx_text("" if value is None else value))
    return (
        f'<c r="{reference}" s="{style}" t="inlineStr">'
        f'<is><t xml:space="preserve">{text}</t></is></c>'
    )


def _excel_column_name(index: int) -> str:
    name = ""
    while index:
        index, remainder = divmod(index - 1, 26)
        name = chr(65 + remainder) + name
    return name


def _xlsx_worksheet(rows: list[tuple[Any, ...]], widths: tuple[int, ...]) -> str:
    column_count = len(rows[0])
    if len(widths) != column_count or any(len(row) != column_count for row in rows):
        raise ValueError("Excel rows and widths must use a consistent column count.")

    columns = tuple(_excel_column_name(index) for index in range(1, column_count + 1))
    row_xml = []
    for row_number, row in enumerate(rows, start=1):
        style = 1 if row_number == 1 else 2 if row_number % 2 == 0 else 0
        cells = "".join(
            _xlsx_cell(f"{column}{row_number}", value, style)
            for column, value in zip(columns, row, strict=True)
        )
        row_xml.append(f'<row r="{row_number}">{cells}</row>')
    last_row = len(rows)
    last_column = columns[-1]
    column_xml = "".join(
        f'<col min="{index}" max="{index}" width="{width}" customWidth="1"/>'
        for index, width in enumerate(widths, start=1)
    )
    return (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
        f'<dimension ref="A1:{last_column}{last_row}"/>'
        '<sheetViews><sheetView workbookViewId="0">'
        '<pane ySplit="1" topLeftCell="A2" activePane="bottomLeft" state="frozen"/>'
        '<selection pane="bottomLeft" activeCell="A2" sqref="A2"/>'
        "</sheetView></sheetViews>"
        '<sheetFormatPr defaultRowHeight="15"/>'
        f"<cols>{column_xml}</cols>"
        f"<sheetData>{''.join(row_xml)}</sheetData>"
        f'<autoFilter ref="A1:{last_column}{last_row}"/>'
        "</worksheet>"
    )


def _build_raw_excel_workbook(
    material_template: list[dict[str, Any]],
    takeoff: dict[str, Any],
    handoff: dict[str, Any],
) -> bytes:
    package_relationships = (
        "http://schemas.openxmlformats.org/package/2006/relationships"
    )
    office_relationships = (
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    )
    spreadsheet_namespace = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    worksheet_content_type = (
        "application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"
    )
    worksheet_relationship = f"{office_relationships}/worksheet"
    workbook = BytesIO()
    sheet_tables = (
        (_material_template_rows(material_template), (28, 58, 48)),
        (_takeoff_rows(takeoff), (28, 48, 15, 15, 15, 15, 12, 18)),
        (_handoff_rows(handoff), (34, 58)),
    )
    with ZipFile(workbook, "w", compression=ZIP_DEFLATED) as archive:
        worksheet_overrides = "".join(
            f'<Override PartName="/xl/worksheets/sheet{index}.xml" '
            f'ContentType="{worksheet_content_type}"/>'
            for index in range(1, 4)
        )
        archive.writestr(
            "[Content_Types].xml",
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            '<Types xmlns="http://schemas.openxmlformats.org/package/2006/'
            'content-types">'
            '<Default Extension="rels" ContentType="application/vnd.'
            'openxmlformats-package.relationships+xml"/>'
            '<Default Extension="xml" ContentType="application/xml"/>'
            '<Override PartName="/xl/workbook.xml" ContentType="application/vnd.'
            'openxmlformats-officedocument.spreadsheetml.sheet.main+xml"/>'
            '<Override PartName="/xl/styles.xml" ContentType="application/vnd.'
            'openxmlformats-officedocument.spreadsheetml.styles+xml"/>'
            f"{worksheet_overrides}"
            "</Types>",
        )
        archive.writestr(
            "_rels/.rels",
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            f'<Relationships xmlns="{package_relationships}">'
            f'<Relationship Id="rId1" Type="{office_relationships}/'
            'officeDocument" Target="xl/workbook.xml"/>'
            "</Relationships>",
        )
        sheets = "".join(
            f'<sheet name="{name}" sheetId="{index}" r:id="rId{index}"/>'
            for index, name in enumerate(_XLSX_SHEETS, start=1)
        )
        archive.writestr(
            "xl/workbook.xml",
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            f'<workbook xmlns="{spreadsheet_namespace}" '
            f'xmlns:r="{office_relationships}">'
            "<bookViews><workbookView/></bookViews>"
            f"<sheets>{sheets}</sheets>"
            '<calcPr calcId="0"/></workbook>',
        )
        relationships = "".join(
            f'<Relationship Id="rId{index}" Type="{worksheet_relationship}" '
            f'Target="worksheets/sheet{index}.xml"/>'
            for index in range(1, 4)
        )
        archive.writestr(
            "xl/_rels/workbook.xml.rels",
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            f'<Relationships xmlns="{package_relationships}">'
            f"{relationships}"
            f'<Relationship Id="rId4" Type="{office_relationships}/styles" '
            'Target="styles.xml"/>'
            "</Relationships>",
        )
        archive.writestr(
            "xl/styles.xml",
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            f'<styleSheet xmlns="{spreadsheet_namespace}">'
            '<fonts count="2"><font><sz val="10"/><name val="Arial"/>'
            '<color rgb="FF000000"/></font><font><b/><sz val="10"/>'
            '<name val="Arial"/><color rgb="FFFFFFFF"/></font></fonts>'
            '<fills count="4"><fill><patternFill patternType="none"/></fill>'
            '<fill><patternFill patternType="gray125"/></fill>'
            '<fill><patternFill patternType="solid"><fgColor rgb="FF595959"/>'
            '<bgColor indexed="64"/></patternFill></fill>'
            '<fill><patternFill patternType="solid"><fgColor rgb="FFF2F2F2"/>'
            '<bgColor indexed="64"/></patternFill></fill></fills>'
            '<borders count="2"><border><left/><right/><top/><bottom/>'
            '<diagonal/></border><border><left style="thin">'
            '<color rgb="FFD9D9D9"/></left><right style="thin">'
            '<color rgb="FFD9D9D9"/></right><top style="thin">'
            '<color rgb="FFD9D9D9"/></top><bottom style="thin">'
            '<color rgb="FFD9D9D9"/></bottom><diagonal/></border></borders>'
            '<cellStyleXfs count="1"><xf numFmtId="0" fontId="0" fillId="0" '
            'borderId="0"/></cellStyleXfs><cellXfs count="3">'
            '<xf numFmtId="0" fontId="0" fillId="0" borderId="1" '
            'applyBorder="1"/><xf numFmtId="0" fontId="1" fillId="2" '
            'borderId="1" applyFont="1" applyFill="1" applyBorder="1" '
            'applyAlignment="1"><alignment vertical="center"/></xf>'
            '<xf numFmtId="0" fontId="0" fillId="3" borderId="1" '
            'applyFill="1" applyBorder="1"/></cellXfs>'
            '<cellStyles count="1"><cellStyle name="Normal" xfId="0" '
            'builtinId="0"/></cellStyles></styleSheet>',
        )
        for index, (rows, widths) in enumerate(sheet_tables, start=1):
            archive.writestr(
                f"xl/worksheets/sheet{index}.xml",
                _xlsx_worksheet(rows, widths),
            )
    return workbook.getvalue()


def _has_display_value(value: Any) -> bool:
    if value is None or isinstance(value, (dict, list, tuple)):
        return False
    if not isinstance(value, (str, int, float, bool)):
        return False
    if not isinstance(value, str):
        return True

    text = value.strip()
    if not text:
        return False
    if text[:1] not in {"{", "["}:
        return True
    try:
        parsed = json.loads(text)
    except json.JSONDecodeError:
        return True
    return not isinstance(parsed, (dict, list))


def _readable_label(value: str) -> str:
    label = value.strip().strip("_").replace("_", " ").replace("-", " ")
    label = re.sub(r"([a-z0-9])([A-Z])", r"\1 \2", label)
    label = re.sub(r"([A-Z]+)([A-Z][a-z])", r"\1 \2", label)
    label = re.sub(r"\bA(\d+)to\s*A(\d+)\b", r"A\1 to A\2", label)
    label = re.sub(r"\s+", " ", label).strip()
    label = re.sub(r"\bbhom\b", "BHoM", label, flags=re.IGNORECASE)
    label = re.sub(r"\bepd\b", "EPD", label, flags=re.IGNORECASE)
    label = re.sub(r"\bguid\b", "GUID", label, flags=re.IGNORECASE)
    return label[:1].upper() + label[1:] if label else "Property"


def _field_label_and_unit(field_name: str) -> tuple[str, str]:
    lower_name = field_name.casefold()
    for encoded_suffix, unit in _UNIT_SUFFIXES:
        if lower_name.endswith(encoded_suffix):
            return _readable_label(field_name[: -len(encoded_suffix)]), unit
    return _readable_label(field_name), ""


def _short_bhom_type(value: Any, *, fallback: str) -> str:
    name = str(value or "").rsplit(".", 1)[-1]
    return _readable_label(name or fallback)


def _dynamic_template_values(
    template: dict[str, Any],
    epd: dict[str, Any],
) -> list[tuple[str, Any, str]]:
    values: list[tuple[str, Any, str]] = []
    label_counts: dict[str, int] = {}

    def append(prefix: str, field_name: str, value: Any) -> None:
        if field_name in _DYNAMIC_EXCLUDED_KEYS or not _has_display_value(value):
            return
        field_label, unit = _field_label_and_unit(field_name)
        label = f"{prefix} · {field_label}" if prefix else field_label
        count = label_counts.get(label, 0) + 1
        label_counts[label] = count
        if count > 1:
            label = f"{label} ({count})"
        values.append((label, value, unit))

    for key in sorted(template, key=str.casefold):
        if key != "Properties":
            append("Material", key, template[key])

    for key in sorted(epd, key=str.casefold):
        if key not in _DYNAMIC_EPD_CONTAINERS:
            append("EPD", key, epd[key])

    for container_name in _DYNAMIC_EPD_CONTAINERS:
        container = epd.get(container_name)
        if isinstance(container, dict):
            container = container.get("_v")
        if not isinstance(container, list):
            continue

        records = [
            (index, record)
            for index, record in enumerate(container)
            if isinstance(record, dict)
        ]
        records.sort(
            key=lambda pair: (
                str(pair[1].get("_t", "")).casefold(),
                str(pair[1].get("Name", "")).casefold(),
                pair[0],
            )
        )
        for _, record in records:
            prefix = _short_bhom_type(
                record.get("_t"),
                fallback=_readable_label(container_name),
            )
            record_items = sorted(
                ((str(key), value) for key, value in record.items()),
                key=lambda item: item[0].casefold(),
            )
            for key, value in record_items:
                if isinstance(value, dict):
                    nested_prefix = f"{prefix} · {_readable_label(key)}"
                    nested_items = sorted(
                        (
                            (str(nested_key), nested_value)
                            for nested_key, nested_value in value.items()
                        ),
                        key=lambda item: item[0].casefold(),
                    )
                    for nested_key, nested_value in nested_items:
                        append(
                            nested_prefix,
                            nested_key,
                            nested_value,
                        )
                else:
                    append(prefix, key, value)

    return values


def _data_item(
    label: str,
    value: Any,
    *,
    suffix: str = "",
    number_of_decimals: int | None = None,
) -> vkt.DataItem:
    if isinstance(value, bool):
        value = "Yes" if value else "No"
    options: dict[str, Any] = {}
    if suffix:
        options["suffix"] = suffix
    if (
        number_of_decimals is not None
        and isinstance(value, (int, float))
        and not isinstance(value, bool)
    ):
        options["number_of_decimals"] = number_of_decimals
    return vkt.DataItem(label, value, **options)


def _excel_material_rows(
    inventory: list[dict[str, Any]],
    mappings: dict[str, str],
    lookup: dict[str, Any],
    templates: list[dict[str, Any]],
) -> list[dict[str, Any]]:
    inventory_by_name = {str(row["material_name"]): row for row in inventory}
    source = catalogue_source(lookup)
    rows: list[dict[str, Any]] = []

    for template in templates:
        source_name = str(template["Name"])
        inventory_row = inventory_by_name[source_name]
        record = record_by_id(lookup, mappings.get(source_name, "")) or {}
        properties = template.get("Properties")
        epd = (
            next(
                (
                    value
                    for value in properties
                    if isinstance(value, dict)
                    and str(value.get("_t", "")).endswith(
                        ".EnvironmentalProductDeclaration"
                    )
                ),
                {},
            )
            if isinstance(properties, list)
            else {}
        )
        dataset_parts = [
            str(value).strip()
            for value in (record.get("dataset_name"), source.get("dataset"))
            if str(value or "").strip()
        ]
        row: dict[str, Any] = {
            "Source material": source_name,
            "Selected material / EPD": epd.get("Name")
            or record.get("name")
            or source_name,
            "Dataset / catalogue": " · ".join(dataset_parts),
            "Material BHoM type": template.get("_t"),
            "EPD BHoM type": epd.get("_t"),
            "BHoM GUID": epd.get("BHoM_Guid") or record.get("catalog_id"),
            "Quantity basis": epd.get("QuantityType") or record.get("quantity_type"),
            "Volume m³": inventory_row.get("volume_m3"),
            "Density kg/m³": template.get("Density"),
            "Mass kg": inventory_row.get("mass_kg"),
        }
        for label, value, unit in _dynamic_template_values(template, epd):
            column = f"{label}{unit}"
            if column not in row:
                row[column] = value
        rows.append(row)
    return rows


def _excel_takeoff_rows(takeoff: dict[str, Any]) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    for index, item in enumerate(takeoff["MaterialTakeoffItems"], start=1):
        material = item["Material"]
        rows.append(
            {
                "Item": index,
                "Material": material.get("Name"),
                "Item BHoM type": item.get("_t"),
                "Material BHoM type": material.get("_t"),
                "Volume m³": item.get("Volume"),
                "Density kg/m³": material_density(material),
                "Mass kg": item.get("Mass"),
                "Area m²": item.get("Area"),
                "Length m": item.get("Length"),
                "Count": item.get("NumberItem"),
            }
        )
    return rows


def _excel_handoff_rows(handoff: dict[str, Any]) -> list[dict[str, Any]]:
    handoff_params = handoff["params"]
    takeoff = json.loads(handoff_params["takeoff_json"])
    templates = json.loads(handoff_params["template_materials_json"])
    catalogue = handoff["source"]["catalogue"]
    return [
        {
            "Schema version": handoff["schema_version"],
            "Target app": handoff["target"]["app"],
            "Target method": handoff["target"]["method_name"],
            "Source app": handoff["source"]["app"],
            "Catalogue dataset": catalogue["dataset"],
            "BHoM component": catalogue["toolkit"],
            "Library path": catalogue["library_path"],
            "Available datasets": catalogue["dataset_count"],
            "Available EPDs": catalogue["epd_count"],
            "Gross floor area m²": handoff_params["gross_floor_area_m2"],
            "Modules": ", ".join(handoff_params["modules"]),
            "Prioritise template materials": handoff_params[
                "prioritise_template_materials"
            ],
            "Template materials": len(templates),
            "Takeoff contract": takeoff["_t"],
            "Takeoff name": takeoff.get("Name"),
            "Takeoff items": len(takeoff["MaterialTakeoffItems"]),
        }
    ]


def _write_excel_sheet(
    workbook: Workbook,
    *,
    sheet_name: str,
    title: str,
    rows: list[dict[str, Any]],
    preferred_columns: tuple[str, ...],
) -> None:
    worksheet = workbook.create_sheet(sheet_name)
    columns = list(preferred_columns)
    columns.extend(
        sorted(
            {
                column
                for row in rows
                for column in row
                if column not in preferred_columns
            },
            key=str.casefold,
        )
    )
    last_column = get_column_letter(len(columns))

    worksheet.append([title])
    worksheet.merge_cells(f"A1:{last_column}1")
    title_cell = worksheet["A1"]
    title_cell.fill = PatternFill("solid", fgColor="FF000000")
    title_cell.font = Font(color="FFFFFFFF", bold=True, size=12)
    title_cell.alignment = Alignment(vertical="center")
    worksheet.row_dimensions[1].height = 24

    worksheet.append(columns)
    header_fill = PatternFill("solid", fgColor="FFD9D9D9")
    border = Border(bottom=Side(style="thin", color="FF808080"))
    for cell in worksheet[2]:
        cell.fill = header_fill
        cell.font = Font(color="FF000000", bold=True)
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        cell.border = border
    worksheet.row_dimensions[2].height = 30

    for row in rows:
        worksheet.append([row.get(column) for column in columns])
    for row in worksheet.iter_rows(min_row=3):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = border
            if isinstance(cell.value, float):
                cell.number_format = "0.00##"
            elif isinstance(cell.value, int) and not isinstance(cell.value, bool):
                cell.number_format = "0"

    worksheet.freeze_panes = "A3"
    worksheet.auto_filter.ref = f"A2:{last_column}{worksheet.max_row}"
    worksheet.sheet_view.showGridLines = False

    for column_index, column in enumerate(columns, start=1):
        values = [column]
        values.extend(
            str(worksheet.cell(row=row, column=column_index).value or "")
            for row in range(3, worksheet.max_row + 1)
        )
        width = min(max(max(len(value) for value in values) + 2, 12), 42)
        worksheet.column_dimensions[get_column_letter(column_index)].width = width


def _build_excel_workbook(
    inventory: list[dict[str, Any]],
    takeoff: dict[str, Any],
    mappings: dict[str, str],
    lookup: dict[str, Any],
    templates: list[dict[str, Any]],
    handoff: dict[str, Any],
) -> bytes:
    workbook = Workbook()
    workbook.remove(workbook.active)
    _write_excel_sheet(
        workbook,
        sheet_name="Material template",
        title="BHoM material template",
        rows=_excel_material_rows(inventory, mappings, lookup, templates),
        preferred_columns=(
            "Source material",
            "Selected material / EPD",
            "Dataset / catalogue",
            "Material BHoM type",
            "EPD BHoM type",
            "BHoM GUID",
            "Quantity basis",
            "Volume m³",
            "Density kg/m³",
            "Mass kg",
        ),
    )
    _write_excel_sheet(
        workbook,
        sheet_name="Takeoff",
        title=str(takeoff.get("Name") or "Normalized BHoM takeoff"),
        rows=_excel_takeoff_rows(takeoff),
        preferred_columns=(
            "Item",
            "Material",
            "Item BHoM type",
            "Material BHoM type",
            "Volume m³",
            "Density kg/m³",
            "Mass kg",
            "Area m²",
            "Length m",
            "Count",
        ),
    )
    _write_excel_sheet(
        workbook,
        sheet_name="Workflow handoff",
        title="Workflow handoff",
        rows=_excel_handoff_rows(handoff),
        preferred_columns=(
            "Schema version",
            "Target app",
            "Target method",
            "Source app",
            "Catalogue dataset",
            "BHoM component",
            "Library path",
            "Available datasets",
            "Available EPDs",
            "Gross floor area m²",
            "Modules",
            "Prioritise template materials",
            "Template materials",
            "Takeoff contract",
            "Takeoff name",
            "Takeoff items",
        ),
    )
    stream = BytesIO()
    workbook.save(stream)
    return stream.getvalue()


class Controller(vkt.Controller):
    label = "BHoM Material Template Mapping"
    parametrization = Parametrization

    @staticmethod
    def _context(
        params: Any,
    ) -> tuple[
        list[dict[str, Any]],
        dict[str, Any],
        dict[str, str],
        dict[str, Any],
    ]:
        try:
            inventory, takeoff = resolve_inputs(params)
            mappings = parse_mapping_state(value_of(params, "mapping_state_json", ""))
            lookup = parse_lookup_results(value_of(params, "lookup_results_json", ""))
            return inventory, takeoff, mappings, lookup
        except (ContractError, UnicodeDecodeError, ValueError) as error:
            raise vkt.UserError(str(error)) from error

    @staticmethod
    def _complete_templates(
        inventory: list[dict[str, Any]],
        mappings: dict[str, str],
        lookup: dict[str, Any],
    ) -> list[dict[str, Any]]:
        missing = [
            str(row["material_name"])
            for row in inventory
            if not mappings.get(str(row["material_name"]))
        ]
        if missing:
            preview = ", ".join(missing[:5])
            if len(missing) > 5:
                preview += f", and {len(missing) - 5} more"
            raise vkt.UserError(
                "Map every source material before exporting the BHoM template. "
                f"Still unmapped: {preview}."
            )
        try:
            return build_template_materials(inventory, mappings, lookup)
        except ContractError as error:
            raise vkt.UserError(str(error)) from error

    def _template_context(
        self,
        params: Any,
    ) -> tuple[
        list[dict[str, Any]],
        dict[str, Any],
        dict[str, str],
        dict[str, Any],
        list[dict[str, Any]],
    ]:
        inventory, takeoff, mappings, lookup = self._context(params)
        templates = self._complete_templates(inventory, mappings, lookup)
        return inventory, takeoff, mappings, lookup, templates

    def _lca_handoff(self, params: Any) -> dict[str, Any]:
        _, takeoff, _, lookup, templates = self._template_context(params)
        return build_handoff(
            takeoff=takeoff,
            template_materials=templates,
            catalogue=catalogue_source(lookup),
            gross_floor_area_m2=float(value_of(params, "gross_floor_area_m2", 0) or 0),
        )

    @staticmethod
    def _search_queries(
        params: Any,
        inventory: list[dict[str, Any]],
    ) -> list[dict[str, Any]]:
        configured: dict[str, str] = {}
        for row in value_of(params, "material_inventory", []) or []:
            name = str(value_of(row, "material_name", "") or "").strip()
            query = str(value_of(row, "search_query", "") or "").strip()
            if name:
                configured[name.casefold()] = query
        return [
            {
                "source_material": str(row["material_name"]),
                "query": configured.get(str(row["material_name"]).casefold())
                or str(row["material_name"]),
                "count": 8,
            }
            for row in inventory
        ]

    def search_bhom_database(self, params: Any, **kwargs: Any):
        inventory, _, _, _ = self._context(params)
        request = {
            "schema_version": "1.0",
            "job_id": str(uuid.uuid4()),
            "dataset_scope": str(
                value_of(params, "dataset_scope", "All installed LCA datasets")
            ),
            "searches": self._search_queries(params, inventory),
        }
        try:
            outputs = execute_worker(
                request_json=json.dumps(request, separators=(",", ":")),
                timeout_seconds=180,
                cache_version=CACHE_VERSION,
            )
            lookup = parse_lookup_results(outputs["lookup-result.json"])
        except Exception as error:
            raise vkt.UserError(
                f"Installed BHoM dataset search failed: {error}"
            ) from error
        if lookup.get("status") != "completed":
            raise vkt.UserError("The BHoM dataset worker did not complete the search.")
        return vkt.SetParamsResult(
            {
                "lookup_results_json": outputs["lookup-result.json"],
                "mapping_state_json": '{"schema_version":"1.0","mappings":{}}',
                "template_materials_json": "[]",
            }
        )

    def finalize_mapping(self, params: Any, **kwargs: Any):
        _, _, _, _, templates = self._template_context(params)
        return vkt.SetParamsResult(
            {
                "template_materials_json": json.dumps(
                    templates,
                    separators=(",", ":"),
                    ensure_ascii=False,
                )
            }
        )

    @vkt.WebView(
        "Mapping",
        duration_guess=1,
        description="Choose one installed BHoM dataset result per material.",
    )
    def mapping_view(self, params: Any, **kwargs: Any):
        inventory, _, mappings, lookup = self._context(params)
        try:
            return vkt.WebResult(html=build_mapping_html(inventory, mappings, lookup))
        except ContractError as error:
            raise vkt.UserError(str(error)) from error

    @vkt.DataView(
        "Material template",
        duration_guess=1,
        description=(
            "Reusable standalone BHoM Material[] built from the saved dataset "
            "selections. This is the same validated template source used by "
            "Export Excel workbook."
        ),
    )
    def template_view(self, params: Any, **kwargs: Any):
        inventory, _, mappings, lookup, templates = self._template_context(params)
        source = catalogue_source(lookup)
        inventory_by_name = {str(row["material_name"]): row for row in inventory}

        summary_items = [
            vkt.DataItem("BHoM contract", "Material[]"),
            vkt.DataItem("Mapped source materials", len(templates)),
            vkt.DataItem("Catalogue", source["dataset"]),
        ]
        material_specs: list[dict[str, Any]] = []

        for template in templates:
            source_name = str(template["Name"])
            row = inventory_by_name[source_name]
            record = record_by_id(lookup, mappings.get(source_name, ""))
            properties = template.get("Properties")
            epd = (
                next(
                    (
                        value
                        for value in properties
                        if isinstance(value, dict)
                        and str(value.get("_t", "")).endswith(
                            ".EnvironmentalProductDeclaration"
                        )
                    ),
                    {},
                )
                if isinstance(properties, list)
                else {}
            )
            epd_name = str(epd.get("Name") or (record or {}).get("name") or source_name)

            dataset_parts = [
                str(value).strip()
                for value in (
                    (record or {}).get("dataset_name"),
                    source.get("dataset"),
                )
                if str(value or "").strip()
            ]
            core_items = [
                vkt.DataItem("Source material", source_name),
                vkt.DataItem("Selected material / EPD", epd_name),
                vkt.DataItem("Dataset / catalogue", " · ".join(dataset_parts)),
                vkt.DataItem(
                    "BHoM type",
                    epd.get("_t") or template["_t"],
                ),
                _data_item(
                    "BHoM GUID",
                    epd.get("BHoM_Guid") or (record or {}).get("catalog_id"),
                ),
                _data_item(
                    "Quantity basis",
                    epd.get("QuantityType") or (record or {}).get("quantity_type"),
                ),
            ]
            core_items = [item for item in core_items if _has_display_value(item.value)]
            for label, key, suffix, decimals in (
                ("Volume", "volume_m3", " m³", 3),
                ("Density", "density_kg_m3", " kg/m³", 1),
                ("Mass", "mass_kg", " kg", 1),
            ):
                value = (
                    template.get("Density") if key == "density_kg_m3" else row.get(key)
                )
                if _has_display_value(value):
                    core_items.append(
                        _data_item(
                            label,
                            value,
                            suffix=suffix,
                            number_of_decimals=decimals,
                        )
                    )

            material_specs.append(
                {
                    "source_name": source_name,
                    "epd_name": epd_name,
                    "core_items": core_items,
                    "extra_values": _dynamic_template_values(template, epd),
                    "selected_extras": [],
                }
            )

        base_item_count = (
            1
            + len(summary_items)
            + sum(1 + len(spec["core_items"]) for spec in material_specs)
        )
        if base_item_count > _DATA_VIEW_ITEM_LIMIT:
            raise vkt.UserError(
                "The Material template result exceeds the VIKTOR DataView "
                "100-item limit before optional properties can be shown. "
                "Use Export Excel to review all flattened scalar fields."
            )

        extra_budget = min(
            _MAX_EXTRA_ITEMS_OVERALL,
            _DATA_VIEW_ITEM_LIMIT - base_item_count,
        )
        for extra_index in range(_MAX_EXTRA_ITEMS_PER_MATERIAL):
            for spec in material_specs:
                extra_values = spec["extra_values"]
                if extra_index >= len(extra_values):
                    continue
                if extra_budget <= 0:
                    break
                spec["selected_extras"].append(extra_values[extra_index])
                extra_budget -= 1
            if extra_budget <= 0:
                break

        items: list[vkt.DataItem] = [
            vkt.DataItem(
                "Material template",
                "Ready",
                status=vkt.DataStatus.SUCCESS,
                status_message=(
                    "Reusable standalone result built from the validated BHoM "
                    "template objects used by the Excel export."
                ),
                subgroup=vkt.DataGroup(*summary_items),
            )
        ]
        for spec in material_specs:
            material_items = list(spec["core_items"])
            material_items.extend(
                _data_item(label, value, suffix=suffix)
                for label, value, suffix in spec["selected_extras"]
            )
            items.append(
                vkt.DataItem(
                    spec["source_name"],
                    spec["epd_name"],
                    status=vkt.DataStatus.SUCCESS,
                    status_message="Mapped to a validated installed BHoM EPD.",
                    subgroup=vkt.DataGroup(*material_items),
                )
            )

        return vkt.DataResult(vkt.DataGroup(*items))

    @vkt.DataView(
        "Workflow handoff",
        duration_guess=1,
        description=(
            "Optional integration envelope that packages the material template, "
            "takeoff, and gross floor area for downstream workflows."
        ),
    )
    def workflow_handoff_view(self, params: Any, **kwargs: Any):
        handoff = self._lca_handoff(params)
        handoff_params = handoff["params"]
        takeoff: dict[str, Any] = json.loads(handoff_params["takeoff_json"])
        templates: list[dict[str, Any]] = json.loads(
            handoff_params["template_materials_json"]
        )
        target = handoff["target"]
        source = handoff["source"]
        catalogue = source["catalogue"]

        items = [
            vkt.DataItem(
                "Workflow handoff",
                "Ready",
                status=vkt.DataStatus.SUCCESS,
                status_message=(
                    "Optional downstream integration envelope. Material template "
                    "is the reusable standalone result."
                ),
                subgroup=vkt.DataGroup(
                    vkt.DataItem("Schema version", handoff["schema_version"]),
                    vkt.DataItem("Target app", target["app"]),
                    vkt.DataItem("Target method", target["method_name"]),
                    vkt.DataItem("Source app", source["app"]),
                    vkt.DataItem("Modules", ", ".join(handoff_params["modules"])),
                ),
            ),
            vkt.DataItem(
                "Packaged inputs",
                "Template + takeoff + area",
                subgroup=vkt.DataGroup(
                    vkt.DataItem("Template materials", len(templates)),
                    vkt.DataItem(
                        "Takeoff contract",
                        takeoff["_t"],
                    ),
                    vkt.DataItem(
                        "Takeoff items",
                        len(takeoff["MaterialTakeoffItems"]),
                    ),
                    vkt.DataItem(
                        "Gross floor area",
                        float(handoff_params["gross_floor_area_m2"]),
                        suffix=" m²",
                        number_of_decimals=2,
                    ),
                    vkt.DataItem(
                        "Prioritise template materials",
                        "Yes"
                        if handoff_params["prioritise_template_materials"]
                        else "No",
                    ),
                ),
            ),
            vkt.DataItem(
                "Catalogue source",
                catalogue["dataset"],
                subgroup=vkt.DataGroup(
                    vkt.DataItem("BHoM component", catalogue["toolkit"]),
                    vkt.DataItem("Library path", catalogue["library_path"]),
                    vkt.DataItem("Available datasets", catalogue["dataset_count"]),
                    vkt.DataItem("Available EPDs", catalogue["epd_count"]),
                ),
            ),
        ]
        return vkt.DataResult(vkt.DataGroup(*items))

    def download_excel(self, params: Any, **kwargs: Any):
        inventory, takeoff, mappings, lookup, templates = self._template_context(params)
        handoff = build_handoff(
            takeoff=takeoff,
            template_materials=templates,
            catalogue=catalogue_source(lookup),
            gross_floor_area_m2=float(value_of(params, "gross_floor_area_m2", 0) or 0),
        )
        return vkt.DownloadResult(
            file_content=_build_excel_workbook(
                inventory,
                takeoff,
                mappings,
                lookup,
                templates,
                handoff,
            ),
            file_name="bhom-material-mapping.xlsx",
        )

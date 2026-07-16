import base64
import json
import re
import unittest
from copy import deepcopy
from io import BytesIO
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

import viktor as vkt
from openpyxl import load_workbook

from app import Controller
from app.catalog import build_template_materials
from app.contracts import build_handoff

LOOKUP_PATH = Path(__file__).parents[1] / "samples" / "lookup-result.sample.json"
CONCRETE_ID = "7b9515a9-7bab-4e98-8542-3fa179c731e5"
STEEL_ID = "a372a265-fd6e-4445-9c03-29ac36e25682"


def default_inventory():
    return [
        SimpleNamespace(
            material_name="Concrete C30/37",
            search_query="ready mix concrete C30/37",
            volume_m3=32.0,
            density_kg_m3=2400.0,
        ),
        SimpleNamespace(
            material_name="Structural Steel",
            search_query="structural steel",
            volume_m3=0.4,
            density_kg_m3=7850.0,
        ),
    ]


def params(
    mapping_state_json='{"schema_version":"1.0","mappings":{}}',
    *,
    lookup_results_json=None,
    material_inventory=None,
):
    return SimpleNamespace(
        takeoff_json="",
        takeoff_file=None,
        dataset_scope="All installed LCA datasets",
        material_inventory=material_inventory or default_inventory(),
        lookup_results_json=(
            lookup_results_json
            if lookup_results_json is not None
            else LOOKUP_PATH.read_text(encoding="utf-8")
        ),
        mapping_state_json=mapping_state_json,
        template_materials_json="[]",
        gross_floor_area_m2=1000,
    )


def call_view(controller: Controller, name: str, view_params: SimpleNamespace):
    decorated_method = Controller.__dict__[name]
    view = decorated_method.__self__
    return view._view_function(controller, params=view_params)


def bootstrap_from_html(html: str):
    match = re.search(r'atob\("([^"]+)"\)', html)
    assert match is not None
    return json.loads(base64.b64decode(match.group(1)).decode("utf-8"))


def workbook_sheet_names(content: bytes):
    return load_workbook(BytesIO(content), read_only=True).sheetnames


def workbook_records(content: bytes, sheet_name: str):
    workbook = load_workbook(BytesIO(content), data_only=True)
    worksheet = workbook[sheet_name]
    headers = [cell.value for cell in worksheet[2]]
    return [
        dict(zip(headers, values, strict=True))
        for values in worksheet.iter_rows(min_row=3, values_only=True)
    ]


def material_sheet_values(content: bytes, material_name: str):
    return next(
        row
        for row in workbook_records(content, "Material template")
        if row["Source material"] == material_name
    )


def complete_mapping_state():
    return json.dumps(
        {
            "schema_version": "1.0",
            "mappings": {
                "Concrete C30/37": CONCRETE_ID,
                "Structural Steel": STEEL_ID,
            },
        }
    )


def concrete_mapping_state():
    return json.dumps(
        {
            "schema_version": "1.0",
            "mappings": {"Concrete C30/37": CONCRETE_ID},
        }
    )


def items_by_label(group):
    return {item.label: item for item in group.values()}


def count_data_items(group):
    count = 0
    for item in group.values():
        count += 1
        subgroup = getattr(item, "_subgroup", None)
        if subgroup is not None:
            count += count_data_items(subgroup)
    return count


def lookup_with_rich_epd():
    lookup = json.loads(LOOKUP_PATH.read_text(encoding="utf-8"))
    concrete = lookup["searches"][0]["candidates"][0]
    epd = json.loads(concrete["epd_json"])
    epd.update(
        {
            "_bhomVersion": "9.2",
            "CustomImpact_kgco2e_per_declared_unit": 1.23,
            "NullProperty": None,
            "EmptyText": " ",
            "EmptyList": [],
            "RawJson": '{"hidden":true}',
            "Fragments": [
                {
                    "_t": ("BH.oM.LifeCycleAssessment.Fragments.AdditionalEPDData"),
                    "Description": ("The compressive strength class follows EN 206."),
                    "LifeSpan": 50,
                    "Manufacturer": "",
                    "ReferenceYear": 2021,
                    "IndustryStandards": [],
                    "Nested": {"TooDeep": {"Hidden": "value"}},
                }
            ],
        }
    )
    epd["EnvironmentalMetrics"][0]["Indicators"]["A4"] = None
    concrete["epd_json"] = json.dumps(epd)
    return lookup


def lookup_with_many_properties(material_count=4, properties_per_material=50):
    source_lookup = json.loads(LOOKUP_PATH.read_text(encoding="utf-8"))
    source_search = source_lookup["searches"][0]
    source_candidate = source_search["candidates"][0]
    searches = []
    inventory = []
    mappings = {}

    for index in range(material_count):
        material_name = f"Material {index + 1}"
        catalog_id = f"catalog-{index + 1}"
        candidate = deepcopy(source_candidate)
        candidate["catalog_id"] = catalog_id
        candidate["name"] = f"Selected EPD {index + 1}"
        epd = json.loads(candidate["epd_json"])
        epd["BHoM_Guid"] = catalog_id
        epd["Name"] = candidate["name"]
        for property_index in range(properties_per_material):
            epd[f"ExtraProperty{property_index:03d}"] = property_index
        candidate["epd_json"] = json.dumps(epd)
        searches.append(
            {
                "source_material": material_name,
                "query": material_name,
                "candidates": [candidate],
            }
        )
        inventory.append(
            SimpleNamespace(
                material_name=material_name,
                search_query=material_name,
                volume_m3=1.0,
                density_kg_m3=1000.0,
            )
        )
        mappings[material_name] = catalog_id

    source_lookup["searches"] = searches
    state = json.dumps({"schema_version": "1.0", "mappings": mappings})
    return source_lookup, inventory, state


class TestViews(unittest.TestCase):
    @patch.dict("os.environ", {"VIKTOR_JS_SDK_PATH": "/sdk/"})
    def test_mapping_view_uses_current_hidden_mapping_state(self):
        state = json.dumps(
            {
                "schema_version": "1.0",
                "mappings": {"Concrete C30/37": CONCRETE_ID},
            }
        )

        result = call_view(Controller(), "mapping_view", params(state))

        self.assertIsInstance(result, vkt.WebResult)
        self.assertIn("Material and EPD mapping", result.html)
        self.assertIn("installed BHoM datasets", result.html)
        self.assertEqual(
            bootstrap_from_html(result.html)["mappings"],
            {"Concrete C30/37": CONCRETE_ID},
        )

    def test_both_result_views_require_complete_mapping(self):
        for view_name in ("template_view", "workflow_handoff_view"):
            with (
                self.subTest(view_name=view_name),
                self.assertRaisesRegex(
                    vkt.UserError,
                    "Still unmapped: Concrete C30/37, Structural Steel",
                ),
            ):
                call_view(Controller(), view_name, params())

    @patch("app.app.build_template_materials", wraps=build_template_materials)
    def test_material_template_view_and_download_share_validated_template_source(
        self,
        template_builder,
    ):
        view_params = params(complete_mapping_state())

        view = call_view(Controller(), "template_view", view_params)
        download = Controller().download_excel(params=view_params)

        self.assertIsInstance(view, vkt.DataResult)
        self.assertEqual(
            Controller.__dict__["template_view"].__self__._label,
            "Material template",
        )
        self.assertEqual(template_builder.call_count, 2)
        self.assertEqual(
            template_builder.call_args_list[0],
            template_builder.call_args_list[1],
        )

        self.assertEqual(download._file_name, "bhom-material-mapping.xlsx")
        template_values = material_sheet_values(
            download._file_content,
            "Concrete C30/37",
        )
        groups = items_by_label(view.data)
        summary = items_by_label(groups["Material template"]._subgroup)
        self.assertEqual(groups["Material template"].value, "Ready")
        self.assertEqual(summary["BHoM contract"].value, "Material[]")
        self.assertEqual(summary["Mapped source materials"].value, 2)

        concrete = items_by_label(groups["Concrete C30/37"]._subgroup)
        self.assertEqual(
            concrete["Selected material / EPD"].value,
            template_values["Selected material / EPD"],
        )
        self.assertEqual(
            concrete["BHoM type"].value,
            template_values["EPD BHoM type"],
        )
        self.assertEqual(
            concrete["BHoM GUID"].value,
            template_values["BHoM GUID"],
        )
        self.assertEqual(
            concrete["Dataset / catalogue"].value,
            "Boverket · All installed LCA datasets",
        )
        self.assertEqual(concrete["Quantity basis"].value, "Mass")
        self.assertEqual(concrete["Volume"].value, 32.0)
        self.assertEqual(concrete["Density"].value, 2400.0)
        self.assertEqual(concrete["Mass"].value, 76_800.0)

    def test_material_template_shows_dynamic_non_empty_scalar_epd_information(self):
        rich_lookup = lookup_with_rich_epd()
        view_params = params(
            concrete_mapping_state(),
            lookup_results_json=json.dumps(rich_lookup),
            material_inventory=[default_inventory()[0]],
        )

        result = call_view(Controller(), "template_view", view_params)

        concrete = items_by_label(
            items_by_label(result.data)["Concrete C30/37"]._subgroup
        )
        self.assertEqual(concrete["EPD · Type"].value, "Sector")
        self.assertEqual(concrete["EPD · BHoM Version"].value, "9.2")
        self.assertEqual(concrete["EPD · Custom Impact"].value, 1.23)
        self.assertEqual(
            concrete["EPD · Custom Impact"].suffix,
            " kgCO₂e/declared unit",
        )
        self.assertEqual(
            concrete["Climate Change Total Metric · Indicators · A1 to A3"].value,
            0.1446,
        )
        self.assertEqual(
            concrete["Additional EPD Data · Description"].value,
            "The compressive strength class follows EN 206.",
        )
        self.assertEqual(
            concrete["Additional EPD Data · Life Span"].value,
            50,
        )
        self.assertEqual(
            concrete["Additional EPD Data · Reference Year"].value,
            2021,
        )

        labels = set(concrete)
        for hidden_label in (
            "EPD · Null Property",
            "EPD · Empty Text",
            "EPD · Empty List",
            "EPD · Raw Json",
            "Additional EPD Data · Manufacturer",
            "Additional EPD Data · Industry Standards",
            "Additional EPD Data · Nested · Too Deep",
        ):
            self.assertNotIn(hidden_label, labels)
        self.assertTrue(
            all(
                not isinstance(item.value, (dict, list, tuple))
                for item in concrete.values()
            )
        )

    def test_material_template_dynamic_items_respect_per_material_and_overall_caps(
        self,
    ):
        lookup, inventory, mapping_state = lookup_with_many_properties()

        result = call_view(
            Controller(),
            "template_view",
            params(
                mapping_state,
                lookup_results_json=json.dumps(lookup),
                material_inventory=inventory,
            ),
        )

        self.assertLessEqual(count_data_items(result.data), 100)
        groups = items_by_label(result.data)
        core_labels = {
            "Source material",
            "Selected material / EPD",
            "Dataset / catalogue",
            "BHoM type",
            "BHoM GUID",
            "Quantity basis",
            "Volume",
            "Density",
            "Mass",
        }
        total_extras = 0
        for index in range(4):
            material_items = items_by_label(groups[f"Material {index + 1}"]._subgroup)
            extras = set(material_items) - core_labels
            self.assertLessEqual(len(extras), 12)
            total_extras += len(extras)
        self.assertLessEqual(total_extras, 40)

    @patch("app.app.build_handoff", wraps=build_handoff)
    def test_workflow_handoff_view_and_download_share_builder_and_label(
        self,
        handoff_builder,
    ):
        view_params = params(complete_mapping_state())

        view = call_view(Controller(), "workflow_handoff_view", view_params)
        download = Controller().download_excel(params=view_params)

        decorator = Controller.__dict__["workflow_handoff_view"].__self__
        self.assertEqual(decorator._label, "Workflow handoff")
        self.assertIn("Optional integration envelope", decorator._description)
        self.assertIsInstance(view, vkt.DataResult)
        self.assertEqual(handoff_builder.call_count, 2)
        self.assertEqual(
            handoff_builder.call_args_list[0],
            handoff_builder.call_args_list[1],
        )

        handoff_values = workbook_records(
            download._file_content,
            "Workflow handoff",
        )[0]
        self.assertEqual(download._file_name, "bhom-material-mapping.xlsx")
        self.assertEqual(handoff_values["Schema version"], "1.0")
        self.assertEqual(handoff_values["Target method"], "run_analysis")
        groups = items_by_label(view.data)
        envelope = items_by_label(groups["Workflow handoff"]._subgroup)
        packaged = items_by_label(groups["Packaged inputs"]._subgroup)
        self.assertEqual(groups["Workflow handoff"].value, "Ready")
        self.assertEqual(
            envelope["Target method"].value,
            handoff_values["Target method"],
        )
        self.assertEqual(
            packaged["Gross floor area"].value,
            handoff_values["Gross floor area m²"],
        )
        self.assertEqual(packaged["Template materials"].value, 2)

    @patch("app.app.execute_worker")
    def test_successful_search_resets_current_mapping_state(self, execute_worker):
        execute_worker.return_value = {
            "lookup-result.json": LOOKUP_PATH.read_text(encoding="utf-8")
        }
        existing_state = json.dumps(
            {
                "schema_version": "1.0",
                "mappings": {"Concrete C30/37": CONCRETE_ID},
            }
        )

        result = Controller().search_bhom_database(params=params(existing_state))

        self.assertEqual(
            result._data["mapping_state_json"],
            '{"schema_version":"1.0","mappings":{}}',
        )
        self.assertEqual(result._data["template_materials_json"], "[]")

    def test_excel_export_opens_with_expected_sheets_and_typed_values(self):
        result = Controller().download_excel(params=params(complete_mapping_state()))

        workbook = load_workbook(BytesIO(result._file_content), data_only=True)
        self.assertEqual(
            workbook.sheetnames,
            ["Material template", "Takeoff", "Workflow handoff"],
        )
        for worksheet in workbook.worksheets:
            self.assertEqual(worksheet.freeze_panes, "A3")
            self.assertTrue(worksheet.auto_filter.ref.startswith("A2:"))
            self.assertTrue(worksheet["A1"].font.bold)
            self.assertEqual(worksheet["A1"].fill.fgColor.rgb, "FF000000")

        concrete_values = material_sheet_values(
            result._file_content,
            "Concrete C30/37",
        )
        takeoff_values = workbook_records(result._file_content, "Takeoff")[0]
        handoff_values = workbook_records(
            result._file_content,
            "Workflow handoff",
        )[0]
        self.assertEqual(
            concrete_values["Selected material / EPD"],
            "Ready-mix made concrete, buildings C30/37",
        )
        self.assertTrue(
            concrete_values["EPD BHoM type"].endswith("EnvironmentalProductDeclaration")
        )
        self.assertEqual(concrete_values["Density kg/m³"], 2400)
        self.assertIsInstance(concrete_values["Density kg/m³"], int)
        self.assertEqual(takeoff_values["Material"], "Concrete C30/37")
        self.assertEqual(takeoff_values["Volume m³"], 32)
        self.assertIsInstance(takeoff_values["Volume m³"], int)
        self.assertEqual(handoff_values["Gross floor area m²"], 1000)
        self.assertIs(handoff_values["Prioritise template materials"], True)
        self.assertEqual(handoff_values["Takeoff items"], 2)
        self.assertEqual(handoff_values["Template materials"], 2)
        self.assertNotIn("takeoff_json", handoff_values)
        self.assertNotIn("template_materials_json", handoff_values)
        self.assertFalse(
            any(
                isinstance(value, str) and value.lstrip().startswith(("{", "["))
                for value in handoff_values.values()
            )
        )
        self.assertFalse(
            any(
                hasattr(Controller, method)
                for method in (
                    "download_template",
                    "download_takeoff",
                    "download_handoff",
                )
            )
        )
